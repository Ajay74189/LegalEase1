from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def create_docx(document_text: str, document_type: str) -> bytes:

    doc = Document()

    # Default font
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("LEGALEASE")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)

    # Document type
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(document_type.upper())
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()

    # Document content
    for line in document_text.split("\n"):

        line = line.strip()

        if not line:
            doc.add_paragraph()
            continue

        # Detect headings
        if (
            line.isupper()
            and len(line) < 100
        ):
            paragraph = doc.add_paragraph()

            run = paragraph.add_run(line)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:
            paragraph = doc.add_paragraph()

            run = paragraph.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    # Footer
    section = doc.sections[0]

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "LegalEase AI — AI-generated draft. "
        "Please have this document reviewed by a qualified legal professional."
    )

    footer_run.font.size = Pt(8)

    # Save to memory instead of disk
    output = BytesIO()

    doc.save(output)

    output.seek(0)

    return output.getvalue()